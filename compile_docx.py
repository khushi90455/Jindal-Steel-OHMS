import os
import re
import json
import subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TARGET_WORDS_MIN = 19500
TARGET_WORDS_MAX = 21000

def add_toc(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    document.add_page_break()

def parse_inline_formatting(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
        else:
            paragraph.add_run(part)

class DocxBuilder:
    def __init__(self):
        self.doc = Document()
        self.word_count = 0
        self.diagrams = []
        self.diagram_index = 0
        self.setup_styles()
        
    def setup_styles(self):
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        
        # Heading styles
        for i in range(1, 4):
            h_style = self.doc.styles[f'Heading {i}']
            h_font = h_style.font
            h_font.name = 'Times New Roman'
            h_font.color.rgb = RGBColor(0, 0, 0)
    
    def extract_diagrams(self, filepath):
        if not os.path.exists(filepath): return
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        blocks = re.findall(r'```mermaid(.*?)```', content, re.DOTALL)
        for b in blocks:
            self.diagrams.append(b.strip())
            
    def process_markdown(self, filepath, stop_if_full=False):
        if not os.path.exists(filepath): return
        with open(filepath, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        
        in_code = False
        code_lang = ""
        code_block = []
        
        for line in lines:
            if stop_if_full and self.word_count > TARGET_WORDS_MAX:
                break
                
            stripped = line.strip('\n')
            
            if stripped.startswith('```'):
                if in_code:
                    if code_lang == 'mermaid':
                        img_path = f'diagram_images/diagram_{self.diagram_index}.png'
                        if os.path.exists(img_path):
                            p = self.doc.add_paragraph()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            run = p.add_run()
                            run.add_picture(img_path, width=Inches(5.5))
                            caption = self.doc.add_paragraph(f'Figure: Mermaid Diagram {self.diagram_index + 1}')
                            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            caption.runs[0].italic = True
                        self.diagram_index += 1
                    else:
                        code_text = '\\n'.join(code_block)
                        self.word_count += len(code_text.split())
                        p = self.doc.add_paragraph(code_text)
                        p.paragraph_format.left_indent = Inches(0.5)
                        for run in p.runs:
                            run.font.name = 'Courier New'
                            run.font.size = Pt(10)
                    in_code = False
                    code_block = []
                else:
                    in_code = True
                    code_lang = stripped[3:].strip()
                continue
                
            if in_code:
                code_block.append(stripped)
                continue
                
            if stripped.startswith('---'): continue
            
            if stripped.startswith('#'):
                level = len(stripped.split(' ')[0])
                if level > 3: level = 3
                text = stripped[level:].strip()
                if level == 1 and len(self.doc.paragraphs) > 2:
                    self.doc.add_page_break()
                self.doc.add_heading(text, level=level)
                self.word_count += len(text.split())
                continue
                
            if stripped.startswith('- ') or stripped.startswith('* '):
                p = self.doc.add_paragraph(style='List Bullet')
                text = stripped[2:]
                parse_inline_formatting(p, text)
                self.word_count += len(text.split())
                continue
                
            if stripped.startswith('|'):
                p = self.doc.add_paragraph()
                clean_table = stripped.replace('|', ' ').strip()
                if '---' not in clean_table:
                    parse_inline_formatting(p, clean_table)
                    self.word_count += len(clean_table.split())
                continue
                
            if stripped.strip():
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 1.5
                parse_inline_formatting(p, stripped)
                self.word_count += len(stripped.split())

def main():
    builder = DocxBuilder()
    
    files_to_process = [
        'JSPulse_OHMS_Final_Comprehensive_Internship_Report.md',
        'JSPulse_OHMS_Technical_Architecture_Flows.md'
    ]
    
    for f in files_to_process:
        builder.extract_diagrams(f)
        
    with open('diagrams.json', 'w') as f:
        json.dump(builder.diagrams, f)
        
    print("Rendering mermaid diagrams via electron...")
    subprocess.run(["powershell", "-Command", "Start-Process -FilePath '.\\node_modules\\electron\\dist\\electron.exe' -ArgumentList 'render_mermaids.js' -Wait -NoNewWindow"])
    
    # Add TOC (Directly starting as requested)
    builder.doc.add_heading('Table of Contents', level=1)
    add_toc(builder.doc)
    
    print("Building DOCX...")
    for f in files_to_process:
        builder.process_markdown(f)
        
    # Append source code to reach 65-70 pages
    if builder.word_count < TARGET_WORDS_MAX:
        builder.doc.add_page_break()
        builder.doc.add_heading("Appendix: Core Source Code", level=1)
        builder.process_markdown('JSPulse_OHMS_Source_Code_Appendix.md', stop_if_full=True)
        
    out_path = 'JSPulse_OHMS_Internship_Report.docx'
    builder.doc.save(out_path)
    print(f"Saved {out_path} with ~{builder.word_count} words.")

if __name__ == '__main__':
    main()
