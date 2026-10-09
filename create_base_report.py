import os
import copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

dir_path = r'C:\Users\KHUSHI\.gemini\antigravity\brain\52546b67-bb16-4f38-9852-f6f9063e206e'

files = [
    'JSPulse_OHMS_Chapter1.docx',
    'JSPulse_OHMS_Chapter2.docx',
    'JSPulse_OHMS_Chapter3_WithSpaces.docx',
    'JSPulse_OHMS_Chapter4_WithSpaces.docx',
    'JSPulse_OHMS_Chapter5_WithSpaces.docx',
    'JSPulse_OHMS_Chapter6_WithSpaces.docx',
    'JSPulse_OHMS_Chapter7_Final_WithBreaks.docx'
]

def add_toc(document):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after  = Pt(24)
    r = p.add_run("TABLE OF CONTENTS")
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.bold = True

    p_toc = document.add_paragraph()
    run = p_toc.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    document.add_page_break()

def merge_documents():
    master_doc = Document()
    
    # Standard A4 size setup
    section = master_doc.sections[0]
    section.page_width    = Inches(8.27)
    section.page_height   = Inches(11.69)
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.0)
    section.right_margin  = Inches(1.0)
    
    # Cover Page
    p_title = master_doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(120)
    p_title.paragraph_format.space_after  = Pt(12)
    r_title = p_title.add_run("JSPULSE OCCUPATIONAL HEALTH MANAGEMENT SYSTEM")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(20)
    
    p_sub = master_doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(12)
    p_sub.paragraph_format.space_after  = Pt(120)
    r_sub = p_sub.add_run("Comprehensive Internship Report\nJindal Steel & Power, Patratu Plant Clinic Operations")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(14)
    r_sub.italic = True
    
    p_details = master_doc.add_paragraph()
    p_details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_details.paragraph_format.space_before = Pt(100)
    r_details = p_details.add_run("Submitted in partial fulfillment of the requirements for the internship program\nDate: June 2026")
    r_details.font.name = "Times New Roman"
    r_details.font.size = Pt(12)
    
    master_doc.add_page_break()

    # Table of Contents
    add_toc(master_doc)

    body = master_doc.element.body
    
    for idx, f in enumerate(files):
        path = os.path.join(dir_path, f)
        doc = Document(path)
        
        if idx > 0:
            master_doc.add_page_break()
            
        print(f"Merging {f}...")
        
        # Skip Chapter 1's redundant Table of Contents heading
        skip_initial = True if idx == 0 else False
        
        for child in doc.element.body:
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                p_xml_text = "".join(child.itertext()).strip()
                if skip_initial:
                    if "Chapter 1: Introduction and Context" in p_xml_text:
                        skip_initial = False
                    else:
                        continue
                
                new_p = copy.deepcopy(child)
                body.append(new_p)
            elif tag == 'tbl':
                if skip_initial:
                    continue
                new_tbl = copy.deepcopy(child)
                body.append(new_tbl)
            elif tag == 'sectPr':
                continue

    # Header & Footer
    for sec in master_doc.sections:
        sec.different_first_page_header_footer = False
        header = sec.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.text = "JS Pulse OHMS Internship Report"
        p_hdr.runs[0].font.name = "Times New Roman"
        p_hdr.runs[0].font.size = Pt(9)
        p_hdr.runs[0].font.italic = True
        
        footer = sec.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p_ftr.text = ""
        run = p_ftr.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)
        p_ftr.runs[0].font.name = "Times New Roman"
        p_ftr.runs[0].font.size = Pt(10)

    out_name = 'JSPulse_OHMS_Merged_Internship_Report.docx'
    master_doc.save(out_name)
    print(f"Base 7-chapter report saved to {out_name}")

if __name__ == '__main__':
    merge_documents()
