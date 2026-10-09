import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

report_path = 'JSPulse_OHMS_Merged_Internship_Report.docx'
final_path  = 'JSPulse_OHMS_Merged_Internship_Report_Final.docx'

def main():
    if not os.path.exists(report_path):
        print(f"ERROR: {report_path} not found!")
        return

    doc = Document(report_path)
    body = doc.element.body

    # Find the XML element for Chapter 3 using a substring match
    ch3_el = None
    for child in body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            text = "".join(child.itertext()).strip()
            if "Chapter 3: Advanced UML and Architecture Flows" in text:
                pPr = child.get_or_add_pPr()
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None and pStyle.get(qn('w:val')) == 'Heading1':
                    ch3_el = child
                    break

    if ch3_el is None:
        # Fallback to simple text matching
        for child in body:
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                text = "".join(child.itertext()).strip()
                if "Chapter 3: Advanced UML and Architecture Flows" in text:
                    ch3_el = child
                    break

    if ch3_el is None:
        print("ERROR: Chapter 3 heading not found in document body!")
        return

    # Delete all body elements before Chapter 3
    elements_to_delete = []
    for child in body:
        if child == ch3_el:
            break
        elements_to_delete.append(child)

    for el in elements_to_delete:
        body.remove(el)

    print(f"Deleted {len(elements_to_delete)} XML elements from the beginning of the report.")

    # After deletion, the first paragraph in doc.paragraphs is Chapter 3
    p_ch3 = doc.paragraphs[0]
    print(f"First paragraph now: {repr(p_ch3.text)}")

    # Insert Page Break before Chapter 3
    p_break = p_ch3.insert_paragraph_before()
    p_break.paragraph_format.space_before = Pt(0)
    p_break.paragraph_format.space_after  = Pt(0)
    run_break = p_break.add_run()
    run_break.add_break(docx.enum.text.WD_BREAK.PAGE)

    # Insert Table of Contents (Index) Field before the page break
    p_toc_field = p_break.insert_paragraph_before()
    p_toc_field.paragraph_format.space_before = Pt(0)
    p_toc_field.paragraph_format.space_after  = Pt(0)
    run_toc = p_toc_field.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run_toc._r.append(fldChar1)
    run_toc._r.append(instrText)
    run_toc._r.append(fldChar2)
    run_toc._r.append(fldChar3)

    # Insert TOC Heading before the TOC field
    p_toc_head = p_toc_field.insert_paragraph_before()
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_before = Pt(24)
    p_toc_head.paragraph_format.space_after  = Pt(24)
    r_toc_head = p_toc_head.add_run("TABLE OF CONTENTS")
    r_toc_head.font.name = "Times New Roman"
    r_toc_head.font.size = Pt(16)
    r_toc_head.bold = True

    # Save the final document
    doc.save(final_path)
    print(f"Successfully edited in-place and saved to {final_path}")

if __name__ == '__main__':
    main()
