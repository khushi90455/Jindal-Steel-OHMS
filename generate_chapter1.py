import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_toc(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    document.add_page_break()

def setup_styles(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    for i in range(1, 4):
        h_style = doc.styles[f'Heading {i}']
        h_font = h_style.font
        h_font.name = 'Times New Roman'
        h_font.color.rgb = RGBColor(0, 0, 0)
        
def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.add_run(text)
    return p

def main():
    doc = Document()
    setup_styles(doc)
    
    # TOC
    doc.add_heading('Table of Contents', level=1)
    add_toc(doc)
    
    # Chapter 1
    doc.add_heading('Chapter 1: Introduction and Context', level=1)
    
    # 1.1 Abstract
    doc.add_heading('1.1 Abstract & Executive Summary', level=2)
    add_paragraph(doc, "Industrial occupational safety is paramount in heavy manufacturing environments. Under the Factories Act, 1948, industrial clinics must track outpatient visits, occupational accidents, periodic health fitness, and medical inventory. Historically, these tasks rely on physical paper ledgers. This project introduces JS Pulse OHMS (Occupational Health Management System), a comprehensive, offline-first digital healthcare management platform designed explicitly for the Jindal Steel & Power (JSP) Patratu Plant Clinic Operations in Jharkhand.")
    add_paragraph(doc, "The JS Pulse OHMS application revolutionizes the day-to-day workflow of the occupational health center by digitizing the entire patient lifecycle, from the registration desk to the consulting doctor, and finally to the pharmacy dispensation unit. By leveraging modern web technologies such as React 19, Vite, and Electron, the application provides a highly performant, cross-platform desktop experience. The core innovation of the system lies in its robust offline architecture, built atop IndexedDB and managed through a custom wrapper called sheetDB, which bypasses legacy browser storage limits and ensures the clinic remains fully operational even during severe industrial network outages.")
    add_paragraph(doc, "Throughout this internship, a deep focus was placed on developing algorithms for DOM virtualization to handle massive medical spreadsheets without lagging, implementing real-time data validation hooks for pharmaceutical inventory, and creating a seamless Electron IPC (Inter-Process Communication) bridge to facilitate secure, native operating system interactions. The ultimate result is a secure, scalable, and highly responsive clinical management tool that drastically reduces transcription errors, automates compliance reporting, and proactively manages drug expiry lifecycles.")
    doc.add_page_break()
    
    # 1.2 Company Profile
    doc.add_heading('1.2 Company Profile: Jindal Steel & Power', level=2)
    add_paragraph(doc, "Jindal Steel & Power (JSP) is an industrial infrastructure powerhouse in India with a massive footprint in steel manufacturing, power generation, coal mining, and infrastructure projects. As a leading player in the core sector, JSP operates several integrated steel plants across the country. The Patratu manufacturing plant, located in Jharkhand, operates highly specialized heavy machinery environments that require stringent safety and medical protocols.")
    add_paragraph(doc, "The primary operational sectors at the Patratu facility include Blast Furnaces, which function as high-temperature iron-reduction reactors; Steel Melting Shops (SMS), which utilize ladle furnaces to process molten steel; Rolling Mills, where heavy hydraulic rollers shape structural steel beams; and Coke Oven Batteries, serving as massive coal carbonization chambers. Operating these intricate and highly volatile plants involves handling molten metal, heavy material transit, high-voltage electrical equipment, and complex gaseous chemical derivatives. Such an environment mandates strict adherence to industrial health and safety standards.")
    add_paragraph(doc, "JSP is profoundly committed to maintaining the highest safety standards, and its Occupational Health Center (OHC) plays a vital role in ensuring that every employee receives timely medical attention, immediate trauma care in the event of an industrial accident, and routine periodic health checkups as mandated by the Factories Act, 1948. The OHC operates 24/7, servicing thousands of factory workers, administrative staff, and contractors. Ensuring the optimal operation of this medical facility is directly correlated with plant productivity and overall workforce morale.")
    doc.add_page_break()
    
    # 1.3 Objectives
    doc.add_heading('1.3 Internship Objectives & Role', level=2)
    add_paragraph(doc, "The primary objective of this internship was to leverage modern software engineering practices to solve real-world clinical data management challenges at JSP's Patratu facility. The goal was to replace an antiquated, error-prone paper system with a state-of-the-art digital infrastructure tailored to the high-stakes environment of an industrial clinic. My role as an intern involved acting as a full-stack developer and systems architect, tasked with designing, developing, and deploying the JS Pulse OHMS platform from scratch.")
    add_paragraph(doc, "The specific objectives established for the internship included the complete digitization of clinical records, converting physical paper ledgers into an offline-first digital database that guarantees data permanence and instant retrieval. A critical secondary objective was workflow automation, integrating patient logs directly with pharmacy dispensing and inventory systems to eliminate redundant data entry and reduce the administrative burden on nursing staff.")
    add_paragraph(doc, "Furthermore, compliance enhancement was a major target. This involved engineering visual alert algorithms and real-time triggers for expired drugs and overdue health examinations. Because the Patratu plant occasionally experiences network fluctuations, building offline resilience into the application was paramount. I was tasked with creating an architecture capable of running continuously without an active internet connection. Finally, I was responsible for implementing advanced data export and reporting services, providing automated mechanisms to generate one-click Excel and PDF reports for management audits, drastically cutting down the hours previously spent manually compiling statistics.")
    doc.add_page_break()
    
    # 1.4 Problem Statement
    doc.add_heading('1.4 Problem Statement & Project Scope', level=2)
    add_paragraph(doc, "The clinic at the JSP Patratu plant previously utilized manual ledger books to record daily patient visits, medicine distributions, and inventory checks. The daily Outpatient Department (OPD) log required nurses and doctors to physically write down patient details, symptoms, diagnostic codes, and prescribed treatments by hand. Following consultation, the pharmacist had to manually copy the patient's name and the prescribed medicine data from the OPD register into a separate Pharmacy Dispense Book to log what was issued.")
    add_paragraph(doc, "This manual, fragmented process suffered from severe data redundancy and a high risk of transcription errors, which could have serious medical consequences. The process of manually recalculating current drug quantities in Stock Books at the end of each shift was notoriously time-consuming and prone to human arithmetic errors. Most critically, the manual system lacked the ability to proactively track near-expiry medicines, leading to potential inventory wastage, and caused massive delays in preparing monthly statutory compliance reports required by government health inspectors.")
    add_paragraph(doc, "The scope of the JS Pulse OHMS project encompasses the complete lifecycle of clinical management within the JSP Patratu facility. It targets the modernization of the OPD tracking system, the pharmacy inventory management, and the reporting engine. The project bounds include the development of a React-based frontend user interface, the engineering of a custom offline IndexedDB storage layer for instantaneous data access, and the integration of Electron IPC protocols for secure local file system operations (such as automated PDF report generation). The system is scoped to be deployed locally on the clinic's administrative terminals, acting as the centralized nervous system for all medical transactions within the plant.")
    doc.add_page_break()
    
    out_path = 'JSPulse_OHMS_Chapter1.docx'
    doc.save(out_path)
    print(f"Saved {out_path}")

if __name__ == "__main__":
    main()
