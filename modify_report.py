import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

report_path = 'JSPulse_OHMS_Merged_Internship_Report.docx'

def main():
    doc = Document(report_path)
    
    # 1. Find Chapter 3 paragraph
    p_ch3 = None
    for p in doc.paragraphs:
        if p.text.strip() == "Chapter 3: Advanced UML and Architecture Flows" and p.style.name == "Heading 1":
            p_ch3 = p
            break
            
    if not p_ch3:
        print("ERROR: Chapter 3 heading not found!")
        return
        
    print(f"Chapter 3 found. Deleting previous paragraphs...")
    
    # Delete all paragraphs before Chapter 3
    paras_to_delete = []
    for p in doc.paragraphs:
        if p == p_ch3:
            break
        paras_to_delete.append(p)
        
    for p in paras_to_delete:
        p._p.getparent().remove(p._p)
        
    print(f"Deleted {len(paras_to_delete)} paragraphs.")
    
    # 2. Insert Cover Page and Table of Contents before p_ch3
    # We insert them before p_ch3 in reverse order of display so they end up in the correct forward order.
    # Display Order:
    # 1. Title
    # 2. Subtitle
    # 3. Details
    # 4. Page Break
    # 5. TOC Heading
    # 6. TOC Field
    # 7. Page Break
    # 8. Chapter 3
    
    # Let's insert them one by one before p_ch3:
    # First, the page break after TOC field
    p_toc_break = p_ch3.insert_paragraph_before()
    p_toc_break.paragraph_format.space_before = Pt(0)
    p_toc_break.paragraph_format.space_after  = Pt(0)
    p_toc_break.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    
    # Second, the TOC Field paragraph
    p_toc_field = p_toc_break.insert_paragraph_before()
    p_toc_field.paragraph_format.space_before = Pt(0)
    p_toc_field.paragraph_format.space_after  = Pt(0)
    run_toc = p_toc_field.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run_toc._r.append(fldChar1)
    run_toc._r.append(instrText)
    run_toc._r.append(fldChar2)
    run_toc._r.append(fldChar3)
    
    # Third, the TOC Heading
    p_toc_head = p_toc_field.insert_paragraph_before()
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_before = Pt(24)
    p_toc_head.paragraph_format.space_after  = Pt(24)
    r_toc_head = p_toc_head.add_run("TABLE OF CONTENTS")
    r_toc_head.font.name = "Times New Roman"
    r_toc_head.font.size = Pt(16)
    r_toc_head.bold = True
    
    # Fourth, the page break after Cover Page details
    p_cov_break = p_toc_head.insert_paragraph_before()
    p_cov_break.paragraph_format.space_before = Pt(0)
    p_cov_break.paragraph_format.space_after  = Pt(0)
    p_cov_break.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    
    # Fifth, Cover Details
    p_cov_det = p_cov_break.insert_paragraph_before()
    p_cov_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_det.paragraph_format.space_before = Pt(100)
    p_cov_det.paragraph_format.space_after  = Pt(0)
    r_cov_det = p_cov_det.add_run("Submitted in partial fulfillment of the requirements for the internship program\nDate: June 2026")
    r_cov_det.font.name = "Times New Roman"
    r_cov_det.font.size = Pt(12)
    
    # Sixth, Cover Subtitle
    p_cov_sub = p_cov_det.insert_paragraph_before()
    p_cov_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_sub.paragraph_format.space_before = Pt(12)
    p_cov_sub.paragraph_format.space_after  = Pt(120)
    r_cov_sub = p_cov_sub.add_run("Comprehensive Internship Report\nJindal Steel & Power, Patratu Plant Clinic Operations")
    r_cov_sub.font.name = "Times New Roman"
    r_cov_sub.font.size = Pt(14)
    r_cov_sub.italic = True
    
    # Seventh, Cover Title
    p_cov_title = p_cov_sub.insert_paragraph_before()
    p_cov_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_title.paragraph_format.space_before = Pt(120)
    p_cov_title.paragraph_format.space_after  = Pt(12)
    r_cov_title = p_cov_title.add_run("JSPULSE OCCUPATIONAL HEALTH MANAGEMENT SYSTEM")
    r_cov_title.bold = True
    r_cov_title.font.name = "Times New Roman"
    r_cov_title.font.size = Pt(20)

    # Save changes
    doc.save(report_path)
    print("Successfully modified report.")

if __name__ == '__main__':
    main()
